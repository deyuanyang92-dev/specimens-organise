"""table_backend.py — 表后端接口 + XlsxBackend。

2026-10-02 设计（docs/superpowers/specs/2026-10-02-sqlite-truth-excel-mirror-design.md，路线 1
"换底不换壳"）：ExcelStore 的业务逻辑不动，底下读写表格的原语收口到这里，第 2 段再加
SqliteBackend。XlsxBackend 的每个方法体都是从 excel_store.py 原样搬过来的：
  * read_rows      ← _read_plain_rows（流式读、sparse dict、COLUMN_ALIASES 归一）/ _rows_from_sheet
  * replace_many   ← _write_plain_rows（新 Workbook + Sheet1）与 _write_changes_and_summary
                     （load 整本 + _replace_sheet）两条路径
  * append_rows    ← _append_row_incremental / _append_index_row（load + ws.append，失败全量重写兜底）
  * headers        ← _headers；stream_columns ← _stream_columns
  * 所有写盘：tmp → 校验（_verify_workbook_file_can_be_reopened）→ tmp.replace(path)
表的 key = xlsx 文件名；一个文件多张 sheet 时用 "文件名::sheet名"。
本模块只依赖 stdlib + openpyxl，不得 import PyQt5。
"""
from __future__ import annotations

import os
import sys
from pathlib import Path
from typing import Any, Callable, Iterable, Iterator, Protocol

Row = dict[str, Any]
SHEET_SEP = "::"


def split_table_key(key: str) -> tuple[str, str | None]:
    """``"修改记录.xlsx::修改汇总"`` → ``("修改记录.xlsx", "修改汇总")``；无 sheet 时第二项为 None。"""
    file_name, _, sheet = str(key).partition(SHEET_SEP)
    return file_name, (sheet or None)


_OPENPYXL: tuple[Any, Any] | None = None


def _openpyxl():
    """懒加载 openpyxl，并沿用 excel_store._ensure_openpyxl 的 numpy 屏蔽：

    openpyxl 若发现 numpy 已可导入会顺手 import 它（2GB 机器上多占几十 MB）。首次导入时若
    numpy 尚未被别人导入，就临时把 sys.modules["numpy"] 置 None 让 openpyxl 走纯 Python 路径。
    """
    global _OPENPYXL
    if _OPENPYXL is not None:
        return _OPENPYXL
    numpy_module = sys.modules.get("numpy")
    blocked = "numpy" not in sys.modules
    if blocked:
        sys.modules["numpy"] = None  # type: ignore[assignment]
    try:
        from openpyxl import Workbook, load_workbook
    finally:
        if blocked:
            sys.modules.pop("numpy", None)
        elif numpy_module is not None:
            sys.modules["numpy"] = numpy_module
    _OPENPYXL = (Workbook, load_workbook)
    return _OPENPYXL


class TableBackend(Protocol):
    """ExcelStore 之下的表格存储接口。值一律字符串；读出的行是 sparse（只含非空值）。"""

    def exists(self, key: str) -> bool: ...

    def headers(self, key: str) -> list[str]: ...

    def read_rows(self, key: str, fallback_headers: list[str] | None = None) -> list[Row]: ...

    def read_many(self, items: list[tuple[str, list[str] | None]]) -> list[list[Row]]: ...

    def replace_rows(self, key: str, headers: list[str], rows: Iterable[Row]) -> None: ...

    def replace_many(self, items: list[tuple[str, list[str], list[Row]]]) -> None: ...

    def append_rows(self, key: str, headers: list[str], rows: Iterable[Row]) -> None: ...

    def stream_columns(self, key: str, wanted_columns: set[str]) -> Iterator[dict[str, str]]: ...

    def version_token(self, key: str) -> float | int: ...

    def close(self) -> None: ...


# ---------------------------------------------------------------------------
# 模块级 xlsx 读取（store 读"别的工作区"的文件、workspace_tables 都复用）
# ---------------------------------------------------------------------------

def _rows_from_ws(ws: Any, to_string: Callable[[object], str], column_aliases: dict[str, str],
                  fallback_headers: list[str] | None) -> list[Row]:
    # 旧 _read_plain_rows / _rows_from_sheet：流式 iter_rows 解析一行处理一行（2GB 机器不爆内存），
    # sparse dict 只留非空字段，最后把旧列名归一到新列名。
    rows_iter = ws.iter_rows(values_only=True)
    try:
        header_row = next(rows_iter)
    except StopIteration:
        return []
    headers = [to_string(value) for value in header_row]
    if fallback_headers:
        headers = headers or list(fallback_headers)
    data: list[Row] = []
    for raw in rows_iter:
        row: Row = {}
        for idx, header in enumerate(headers):
            if not header or idx >= len(raw):
                continue
            value = to_string(raw[idx])
            if value != "":
                row[header] = value
        if row:
            data.append(row)
    for row in data:
        for old_col, new_col in column_aliases.items():
            if old_col in row and new_col not in row:
                row[new_col] = row.pop(old_col)
    return data


def xlsx_read_rows(path: Path | str, to_string: Callable[[object], str], column_aliases: dict[str, str],
                   fallback_headers: list[str] | None = None, sheet: str | None = None) -> list[Row]:
    path = Path(path)
    if not path.exists():
        return []
    _, load_workbook = _openpyxl()
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet is not None:
            if sheet not in wb.sheetnames:
                return []
            ws = wb[sheet]
        else:
            ws = wb.active
        return _rows_from_ws(ws, to_string, column_aliases, fallback_headers)
    finally:
        wb.close()


def xlsx_headers(path: Path | str, to_string: Callable[[object], str], sheet: str | None = None) -> list[str]:
    path = Path(path)
    if not path.exists():
        return []
    _, load_workbook = _openpyxl()
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet is not None:
            if sheet not in wb.sheetnames:
                return []
            ws = wb[sheet]
        else:
            ws = wb.active
        try:
            return [to_string(cell.value) for cell in next(ws.iter_rows(max_row=1))]
        except StopIteration:
            return []
    finally:
        wb.close()


# ---------------------------------------------------------------------------
# XlsxBackend
# ---------------------------------------------------------------------------

class XlsxBackend:
    """旧模式：每张表一本 xlsx（或一本里的一张 sheet），整表重写 + 原子替换。"""

    def __init__(
        self,
        data_dir: Path | str,
        *,
        fit_row: Callable[[Row, list[str]], Row],
        to_string: Callable[[object], str],
        verify_file: Callable[[Path], None] | None = None,
        column_aliases: dict[str, str] | None = None,
    ) -> None:
        self.data_dir = Path(data_dir)
        self._fit_row = fit_row
        self._to_string = to_string
        self._verify_file = verify_file
        self._column_aliases = dict(column_aliases or {})

    # -- helpers ------------------------------------------------------------

    def _path(self, key: str) -> Path:
        return self.data_dir / split_table_key(key)[0]

    def _atomic_save(self, wb: Any, path: Path) -> None:
        # 旧：tmp = path.with_suffix(f".{os.getpid()}.tmp"); wb.save(tmp); 校验; tmp.replace(path)
        tmp = path.with_suffix(f".{os.getpid()}.tmp")
        try:
            wb.save(tmp)
            if self._verify_file is not None:
                self._verify_file(tmp)
            tmp.replace(path)
        finally:
            try:
                tmp.unlink(missing_ok=True)
            except OSError:
                pass

    def _fill(self, ws: Any, headers: list[str], rows: Iterable[Row]) -> None:
        ws.append(list(headers))
        for row in rows:
            fitted = self._fit_row(row, headers)
            ws.append([fitted.get(header, "") for header in headers])

    # -- TableBackend -------------------------------------------------------

    def exists(self, key: str) -> bool:
        path = self._path(key)
        if not path.exists():
            return False
        _, sheet = split_table_key(key)
        if sheet is None:
            return True
        _, load_workbook = _openpyxl()
        wb = load_workbook(path, read_only=True)
        try:
            return sheet in wb.sheetnames
        finally:
            wb.close()

    def headers(self, key: str) -> list[str]:
        return xlsx_headers(self._path(key), self._to_string, split_table_key(key)[1])

    def read_rows(self, key: str, fallback_headers: list[str] | None = None) -> list[Row]:
        return xlsx_read_rows(
            self._path(key), self._to_string, self._column_aliases, fallback_headers, split_table_key(key)[1]
        )

    def read_many(self, items: list[tuple[str, list[str] | None]]) -> list[list[Row]]:
        """多张表一次读：同一文件只 load 一次（修改记录两张 sheet 走这里）。结果顺序与 items 一致。"""
        results: list[list[Row] | None] = [None] * len(items)
        by_file: dict[str, list[int]] = {}
        for idx, (key, _) in enumerate(items):
            by_file.setdefault(split_table_key(key)[0], []).append(idx)
        _, load_workbook = _openpyxl()
        for file_name, indices in by_file.items():
            path = self.data_dir / file_name
            if not path.exists():
                for idx in indices:
                    results[idx] = []
                continue
            wb = load_workbook(path, read_only=True, data_only=True)
            try:
                for idx in indices:
                    key, fallback = items[idx]
                    _, sheet = split_table_key(key)
                    if sheet is None:
                        ws = wb.active
                    elif sheet in wb.sheetnames:
                        ws = wb[sheet]
                    else:
                        results[idx] = []
                        continue
                    results[idx] = _rows_from_ws(ws, self._to_string, self._column_aliases, fallback)
            finally:
                wb.close()
        return [r if r is not None else [] for r in results]

    def version_token(self, key: str) -> float:
        # 旧 _cached_rows / _ensure_index_voucher_set 等：用文件 mtime 判缓存失效；缺文件按 0.0
        try:
            return self._path(key).stat().st_mtime
        except OSError:
            return 0.0

    def replace_rows(self, key: str, headers: list[str], rows: Iterable[Row]) -> None:
        self.replace_many([(key, list(headers), list(rows))])

    def replace_many(self, items: list[tuple[str, list[str], list[Row]]]) -> None:
        Workbook, load_workbook = _openpyxl()
        by_file: dict[str, list[tuple[str | None, list[str], list[Row]]]] = {}
        for key, headers, rows in items:
            file_name, sheet = split_table_key(key)
            by_file.setdefault(file_name, []).append((sheet, list(headers), list(rows)))
        for file_name, sheet_items in by_file.items():
            path = self.data_dir / file_name
            if all(sheet is None for sheet, _, _ in sheet_items):
                # 旧 _write_plain_rows：新建 Workbook，Sheet1，表头 + 行，原子替换
                sheet, headers, rows = sheet_items[-1]
                wb = Workbook()
                try:
                    ws = wb.active
                    ws.title = "Sheet1"
                    self._fill(ws, headers, rows)
                    self._atomic_save(wb, path)
                finally:
                    try:
                        wb.close()
                    except Exception:
                        pass
                continue
            # 旧 _write_changes_and_summary / 修改明细整表替换：load 整本，_replace_sheet
            # （delete_rows + 表头 + 行），其它 sheet 原样保留，原子替换。
            # 优化：先用 read_only 看一眼 sheet 名（毫秒级）；文件里所有 sheet 都在本次替换范围内时
            # 直接新建 Workbook（省掉整本 load，1500 条标本的修改记录 ≈ 0.1 s/次）；
            # 有不认识的 sheet（用户自加）才走旧的整本 load 以保留它。
            replacing = {sheet or "Sheet1" for sheet, _, _ in sheet_items}
            existing_sheets: list[str] = []
            if path.exists():
                probe = load_workbook(path, read_only=True)
                try:
                    existing_sheets = list(probe.sheetnames)
                finally:
                    probe.close()
            if existing_sheets and not set(existing_sheets).issubset(replacing):
                wb = load_workbook(path)
            else:
                wb = Workbook()
                wb.remove(wb.active)
                # 保持原 sheet 顺序：先按文件里的顺序建，缺的再按本次顺序补
                order = [n for n in existing_sheets if n in replacing] + [n for n in (sheet or "Sheet1" for sheet, _, _ in sheet_items) if n not in existing_sheets]
                for name in order:
                    if name not in wb.sheetnames:
                        wb.create_sheet(name)
            try:
                for sheet, headers, rows in sheet_items:
                    name = sheet or "Sheet1"
                    ws = wb[name] if name in wb.sheetnames else wb.create_sheet(name)
                    if ws.max_row:
                        ws.delete_rows(1, ws.max_row)
                    self._fill(ws, headers, rows)
                self._atomic_save(wb, path)
            finally:
                try:
                    wb.close()
                except Exception:
                    pass

    def append_rows(self, key: str, headers: list[str], rows: Iterable[Row]) -> None:
        rows = list(rows)
        if not rows:
            return
        headers = list(headers)
        path = self._path(key)
        _, sheet = split_table_key(key)
        if not path.exists():
            self.replace_rows(key, headers, rows)
            return
        _, load_workbook = _openpyxl()
        try:
            # 旧 _append_row_incremental / _append_index_row / _ensure_summary_row：
            # load_workbook + ws.append + 原子替换
            wb = load_workbook(path)
            try:
                if sheet is None:
                    ws = wb.active
                elif sheet in wb.sheetnames:
                    ws = wb[sheet]
                else:
                    ws = wb.create_sheet(sheet)
                    ws.append(headers)
                for row in rows:
                    fitted = self._fit_row(row, headers)
                    ws.append([fitted.get(header, "") for header in headers])
                self._atomic_save(wb, path)
            finally:
                try:
                    wb.close()
                except Exception:
                    pass
        except Exception:
            # 旧兜底：load/save 异常 → 全量重写
            existing = self.read_rows(key, headers)
            self.replace_rows(key, headers, existing + rows)

    def stream_columns(self, key: str, wanted_columns: set[str]) -> Iterator[dict[str, str]]:
        # 旧 _stream_columns：流式只取 wanted 列，sparse dict；wanted 含新列名时也匹配旧列名并归一为新名
        path = self._path(key)
        if not path.exists():
            return
        _, load_workbook = _openpyxl()
        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            ws = wb.active
            rows_iter = ws.iter_rows(values_only=True)
            try:
                header_row = next(rows_iter)
            except StopIteration:
                return
            headers = [self._to_string(value) for value in header_row]
            alias_reverse = {new: old for old, new in self._column_aliases.items()}
            extra_wanted = {alias_reverse[h] for h in wanted_columns if h in alias_reverse}
            effective_wanted = set(wanted_columns) | extra_wanted
            wanted_idx = [
                (i, self._column_aliases.get(h, h)) for i, h in enumerate(headers) if h in effective_wanted
            ]
            for raw in rows_iter:
                row: dict[str, str] = {}
                for idx, canonical in wanted_idx:
                    if idx < len(raw):
                        value = self._to_string(raw[idx])
                        if value != "":
                            row[canonical] = value
                if row:
                    yield row
        finally:
            wb.close()

    def close(self) -> None:
        return None
