"""手动输入的新物种保存后记入工作区 字段模版/表格信息预设字段.xlsx，下次自动匹配（2026-10-02 用户需求）。"""
import shutil
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

from specimen_app.species import PRESET_HEADERS, SpeciesMatcher, remember_species


def _write_preset(path: Path, rows):
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook(); ws = wb.active
    ws.append(PRESET_HEADERS)
    for r in rows:
        ws.append(list(r))
    wb.save(path)


class SpeciesMemoryTests(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())
        self.bundled = self.tmp / "bundled" / "表格信息预设字段.xlsx"
        _write_preset(self.bundled, [("中华绒螯蟹", "Eriocheir sinensis", "弓蟹科", "Varunidae")])
        self.ws_preset = self.tmp / "ws" / "字段模版" / "表格信息预设字段.xlsx"

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def test_remember_creates_workspace_preset_and_matcher_finds_it(self):
        m = SpeciesMatcher(self.bundled, extra_paths=[self.ws_preset])
        self.assertIsNone(m.resolve_unique_species("新物种甲"))
        self.assertTrue(remember_species(self.ws_preset, "新物种甲", "Novus alpha", "某科", "Aliidae"))
        self.assertTrue(self.ws_preset.exists())
        hit = m.resolve_unique_species("新物种甲")
        self.assertIsNotNone(hit)
        self.assertEqual(hit.latin_name, "Novus alpha")
        self.assertEqual(hit.family_latin, "Aliidae")
        # 自带库仍可用
        self.assertIsNotNone(m.resolve_unique_species("中华绒螯蟹"))

    def test_remember_appends_and_keeps_existing_rows(self):
        _write_preset(self.ws_preset, [("旧物种", "Vetus a", "旧科", "Vetidae")])
        self.assertTrue(remember_species(self.ws_preset, "新物种乙", "Novus beta", "", ""))
        ws = load_workbook(self.ws_preset).active
        names = [r[0] for r in ws.iter_rows(min_row=2, values_only=True)]
        self.assertEqual(names, ["旧物种", "新物种乙"])
        self.assertEqual([c.value for c in ws[1]][:4], PRESET_HEADERS)

    def test_remember_skips_duplicates_and_blank(self):
        self.assertTrue(remember_species(self.ws_preset, "新物种丙", "Novus gamma", "某科", "Aliidae"))
        self.assertFalse(remember_species(self.ws_preset, "新物种丙", "Novus gamma", "某科", "Aliidae"))
        self.assertFalse(remember_species(self.ws_preset, "  ", "x", "y", "z"))
        ws = load_workbook(self.ws_preset).active
        self.assertEqual(ws.max_row, 2)

    def test_workspace_entry_overrides_bundled_same_chinese_name(self):
        # 自带库写错了拉丁名，用户在工作区更正 → 以工作区为准，且仍能唯一匹配自动填充
        _write_preset(self.ws_preset, [("中华绒螯蟹", "Eriocheir sinensis H. Milne Edwards", "弓蟹科", "Varunidae")])
        m = SpeciesMatcher(self.bundled, extra_paths=[self.ws_preset])
        hit = m.resolve_unique_species("中华绒螯蟹")
        self.assertIsNotNone(hit)
        self.assertEqual(hit.latin_name, "Eriocheir sinensis H. Milne Edwards")

    def test_knows_species(self):
        m = SpeciesMatcher(self.bundled, extra_paths=[self.ws_preset])
        self.assertTrue(m.knows_species("中华绒螯蟹"))
        self.assertFalse(m.knows_species("新物种丁"))

    def test_corrupt_workspace_preset_not_overwritten(self):
        self.ws_preset.parent.mkdir(parents=True)
        self.ws_preset.write_bytes(b"not an xlsx")
        self.assertFalse(remember_species(self.ws_preset, "新物种戊", "N e", "", ""))
        self.assertEqual(self.ws_preset.read_bytes(), b"not an xlsx")  # 不破坏用户文件


class SaveHookTests(unittest.TestCase):
    """分类信息 保存 → _remember_new_species 写工作区预设；已知物种 / 信息不足不写。"""

    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())
        self.ws = self.tmp / "ws"
        self.ws.mkdir()

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def _window(self, values):
        from specimen_app import ui

        class _W:
            def __init__(self, text):
                self._t = text
            def text(self):
                return self._t

        class Fake:
            pass

        w = Fake()
        w.workspace_root = self.ws
        w.class_widgets = {k: _W(v) for k, v in values.items()}
        w.matcher = ui._species_matcher(self.ws)
        return w, ui.SpecimenWindow._remember_new_species

    def test_new_species_is_remembered_then_autofills(self):
        w, fn = self._window({"种名*": "测试新种", "种拉丁": "Testus novus", "科*": "测试科", "科拉丁": "Testidae"})
        self.assertEqual(fn(w), "测试新种")
        self.assertTrue((self.ws / "字段模版" / "表格信息预设字段.xlsx").exists())
        hit = w.matcher.resolve_unique_species("测试新种")
        self.assertIsNotNone(hit)
        self.assertEqual(hit.family_latin, "Testidae")
        self.assertIsNone(fn(w))  # 第二次保存：已知，不重复写

    def test_name_only_is_not_remembered(self):
        w, fn = self._window({"种名*": "只有中文", "种拉丁": "", "科*": "", "科拉丁": ""})
        self.assertIsNone(fn(w))
        self.assertFalse((self.ws / "字段模版" / "表格信息预设字段.xlsx").exists())


if __name__ == "__main__":
    unittest.main()
