"""表后端契约测试（第 1 段：XlsxBackend；第 2 段加 SqliteBackend）。

设计：docs/superpowers/specs/2026-10-02-sqlite-truth-excel-mirror-design.md
计划：docs/superpowers/plans/2026-10-02-sqlite-truth-excel-mirror.md（Task 1 / Task 2）
"""
from __future__ import annotations

import shutil
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from specimen_app.table_backend import SHEET_SEP, XlsxBackend, split_table_key, xlsx_read_rows


def _fit(row, headers):
    return {h: str(row.get(h, "") or "") for h in headers}


def _s(v):
    return "" if v is None else str(v)


class SplitKeyTests(unittest.TestCase):
    def test_split(self):
        self.assertEqual(split_table_key("标本信息.xlsx"), ("标本信息.xlsx", None))
        self.assertEqual(split_table_key("修改记录.xlsx" + SHEET_SEP + "修改汇总"), ("修改记录.xlsx", "修改汇总"))


class XlsxBackendContractTests(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())
        self.data = self.tmp / "数据"
        self.data.mkdir()
        self.b = self.make_backend()

    def make_backend(self):
        return XlsxBackend(
            self.data, fit_row=_fit, to_string=_s, verify_file=None,
            column_aliases={"采集地点缩写*": "采集地缩写*"},
        )

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def test_replace_then_read_is_sparse_and_ordered(self):
        self.b.replace_rows("t.xlsx", ["a", "b"], [{"a": "1", "b": ""}, {"a": "", "b": "2"}])
        self.assertTrue(self.b.exists("t.xlsx"))
        self.assertEqual(self.b.headers("t.xlsx"), ["a", "b"])
        self.assertEqual(self.b.read_rows("t.xlsx"), [{"a": "1"}, {"b": "2"}])

    def test_missing_table_reads_empty(self):
        self.assertFalse(self.b.exists("nope.xlsx"))
        self.assertEqual(self.b.headers("nope.xlsx"), [])
        self.assertEqual(self.b.read_rows("nope.xlsx"), [])
        self.assertEqual(list(self.b.stream_columns("nope.xlsx", {"a"})), [])

    def test_append_rows_appends_in_order(self):
        self.b.replace_rows("t.xlsx", ["a"], [{"a": "1"}])
        self.b.append_rows("t.xlsx", ["a"], [{"a": "2"}, {"a": "3"}])
        self.assertEqual([r["a"] for r in self.b.read_rows("t.xlsx")], ["1", "2", "3"])

    def test_append_creates_missing_file(self):
        self.b.append_rows("n.xlsx", ["a"], [{"a": "x"}])
        self.assertEqual(self.b.read_rows("n.xlsx"), [{"a": "x"}])

    def test_sheet_keys_share_one_workbook(self):
        self.b.replace_many([
            ("c.xlsx" + SHEET_SEP + "修改明细", ["x"], [{"x": "d1"}]),
            ("c.xlsx" + SHEET_SEP + "修改汇总", ["y"], [{"y": "s1"}]),
        ])
        wb = load_workbook(self.data / "c.xlsx", read_only=True)
        self.assertEqual(wb.sheetnames, ["修改明细", "修改汇总"])
        wb.close()
        self.assertTrue(self.b.exists("c.xlsx" + SHEET_SEP + "修改汇总"))
        self.assertFalse(self.b.exists("c.xlsx" + SHEET_SEP + "不存在"))
        self.assertEqual(self.b.read_rows("c.xlsx" + SHEET_SEP + "修改汇总"), [{"y": "s1"}])
        self.b.append_rows("c.xlsx" + SHEET_SEP + "修改汇总", ["y"], [{"y": "s2"}])
        self.assertEqual([r["y"] for r in self.b.read_rows("c.xlsx" + SHEET_SEP + "修改汇总")], ["s1", "s2"])
        self.assertEqual(self.b.read_rows("c.xlsx" + SHEET_SEP + "修改明细"), [{"x": "d1"}])
        # 单 sheet 整表替换不碰另一张
        self.b.replace_rows("c.xlsx" + SHEET_SEP + "修改明细", ["x"], [{"x": "d9"}])
        self.assertEqual(self.b.read_rows("c.xlsx" + SHEET_SEP + "修改明细"), [{"x": "d9"}])
        self.assertEqual([r["y"] for r in self.b.read_rows("c.xlsx" + SHEET_SEP + "修改汇总")], ["s1", "s2"])

    def test_read_many_loads_each_file_once_and_keeps_order(self):
        self.b.replace_many([
            ("c.xlsx" + SHEET_SEP + "修改明细", ["x"], [{"x": "d1"}]),
            ("c.xlsx" + SHEET_SEP + "修改汇总", ["y"], [{"y": "s1"}]),
        ])
        self.b.replace_rows("t.xlsx", ["a"], [{"a": "1"}])
        got = self.b.read_many([
            ("c.xlsx" + SHEET_SEP + "修改汇总", ["y"]),
            ("t.xlsx", None),
            ("c.xlsx" + SHEET_SEP + "修改明细", ["x"]),
            ("missing.xlsx", ["z"]),
        ])
        self.assertEqual(got, [[{"y": "s1"}], [{"a": "1"}], [{"x": "d1"}], []])

    def test_replace_many_preserves_unknown_extra_sheets(self):
        from openpyxl import Workbook
        wb = Workbook(); ws = wb.active; ws.title = "修改明细"; ws.append(["x"]); ws.append(["d1"])
        wb.create_sheet("修改汇总").append(["y"]); extra = wb.create_sheet("用户自加"); extra.append(["keep"]); extra.append(["me"])
        wb.save(self.data / "c.xlsx"); wb.close()
        self.b.replace_many([
            ("c.xlsx" + SHEET_SEP + "修改明细", ["x"], [{"x": "d2"}]),
            ("c.xlsx" + SHEET_SEP + "修改汇总", ["y"], [{"y": "s2"}]),
        ])
        wb = load_workbook(self.data / "c.xlsx", read_only=True)
        self.assertEqual(wb.sheetnames, ["修改明细", "修改汇总", "用户自加"])
        self.assertEqual([r[0] for r in wb["用户自加"].iter_rows(values_only=True)], ["keep", "me"])
        wb.close()
        self.assertEqual(self.b.read_rows("c.xlsx" + SHEET_SEP + "修改明细"), [{"x": "d2"}])

    def test_read_normalises_column_aliases(self):
        self.b.replace_rows("t.xlsx", ["采集地点缩写*"], [{"采集地点缩写*": "QD"}])
        self.assertEqual(self.b.read_rows("t.xlsx"), [{"采集地缩写*": "QD"}])

    def test_version_token_changes_on_write_and_missing_is_zero(self):
        self.assertEqual(self.b.version_token("t.xlsx"), 0)
        self.b.replace_rows("t.xlsx", ["a"], [{"a": "1"}])
        self.assertGreater(self.b.version_token("t.xlsx"), 0)

    def test_stream_columns_only_wanted(self):
        self.b.replace_rows("t.xlsx", ["a", "b"], [{"a": "1", "b": "2"}, {"a": "3", "b": ""}])
        self.assertEqual(list(self.b.stream_columns("t.xlsx", {"b"})), [{"b": "2"}])

    def test_no_tmp_left_behind(self):
        self.b.replace_rows("t.xlsx", ["a"], [{"a": "1"}])
        self.assertEqual([p.name for p in self.data.iterdir()], ["t.xlsx"])

    def test_values_are_strings_even_for_numbers(self):
        self.b.replace_rows("t.xlsx", ["n"], [{"n": 12}])
        self.assertEqual(self.b.read_rows("t.xlsx"), [{"n": "12"}])

    def test_module_reader_reads_external_path(self):
        self.b.replace_rows("t.xlsx", ["a"], [{"a": "1"}])
        self.assertEqual(xlsx_read_rows(self.data / "t.xlsx", _s, {}), [{"a": "1"}])
        self.assertEqual(xlsx_read_rows(self.data / "missing.xlsx", _s, {}), [])


class StoreUsesBackendTests(unittest.TestCase):
    """Task 2：ExcelStore 的读写原语全部经过 self._backend。"""

    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def test_store_routes_reads_and_writes_through_backend(self):
        from specimen_app.excel_store import ExcelStore
        from specimen_app.models import CHANGE_LOG_FILE, SPECIMEN_FILE

        store = ExcelStore(self.tmp)
        try:
            self.assertEqual(store.storage_backend_name, "xlsx")
            v = store.create_specimen()
            store.set_fields("specimen", v, {"备注": "via backend"})
            calls: list[str] = []
            orig = store._backend.read_rows

            def spy(key, fallback_headers=None):
                calls.append(key)
                return orig(key, fallback_headers)

            store._backend.read_rows = spy  # type: ignore[method-assign]
            store._invalidate_cache(SPECIMEN_FILE)
            self.assertEqual(store.get_specimen(v)["备注"], "via backend")
            self.assertIn(SPECIMEN_FILE, calls)
            detail = store._backend.read_rows(CHANGE_LOG_FILE + SHEET_SEP + "修改明细", None)
            self.assertTrue(any(r.get("字段名") == "备注" and r.get("新值") == "via backend" for r in detail))
        finally:
            store.close()

    def test_table_key_for_external_path_is_none(self):
        from specimen_app.excel_store import ExcelStore

        store = ExcelStore(self.tmp)
        try:
            self.assertEqual(store._table_key_for(store.data_dir / "标本信息.xlsx"), "标本信息.xlsx")
            self.assertIsNone(store._table_key_for(self.tmp / "别处" / "标本信息.xlsx"))
            self.assertIsNone(store._table_key_for(store.data_dir / "数据版本" / "x" / "标本信息.xlsx"))
        finally:
            store.close()


if __name__ == "__main__":
    unittest.main()
