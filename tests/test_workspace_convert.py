"""第 2 段：模式识别（Task 5）与旧工作区自动转换（Task 6）。

设计：docs/superpowers/specs/2026-10-02-sqlite-truth-excel-mirror-design.md 第 1–2 节。
"""
from __future__ import annotations

import json
import os
import shutil
import tempfile
import unittest
from pathlib import Path

from specimen_app.excel_store import ExcelStore
from specimen_app.models import SPECIMEN_FILE, SQLITE_DATA_FILE, WORKSPACE_CONFIG_FILE


class ModeSelectionTests(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())
        self.data = self.tmp / "数据"

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def test_sqlite_workspace_basics(self):
        store = ExcelStore(self.tmp, backend="sqlite")
        try:
            self.assertEqual(store.storage_backend_name, "sqlite")
            self.assertTrue((self.data / SQLITE_DATA_FILE).exists())
            self.assertFalse((self.data / SPECIMEN_FILE).exists(), "镜像 xlsx 要到导出时才有")
            self.assertEqual(store.config.get("storage_backend"), "sqlite")
            v = store.create_specimen()
            store.set_fields("specimen", v, {"备注": "sqlite 模式"})
            self.assertEqual(store.get_specimen(v)["备注"], "sqlite 模式")
            self.assertEqual(store.get_specimen(v)["入库编号*"], v)
        finally:
            store.close()

    def test_backend_kwarg_forces_xlsx(self):
        store = ExcelStore(self.tmp, backend="xlsx")
        try:
            self.assertEqual(store.storage_backend_name, "xlsx")
            self.assertFalse((self.data / SQLITE_DATA_FILE).exists())
            self.assertTrue((self.data / SPECIMEN_FILE).exists())
            self.assertNotIn("storage_backend", store.config)
        finally:
            store.close()

    def test_reopen_detects_sqlite(self):
        store = ExcelStore(self.tmp, backend="sqlite")
        v = store.create_specimen()
        store.close()
        store = ExcelStore(self.tmp)
        try:
            self.assertEqual(store.storage_backend_name, "sqlite")
            self.assertEqual(store.get_specimen(v)["入库编号*"], v)
        finally:
            store.close()

    def test_open_repairs_config_when_sqlite_exists_without_flag(self):
        store = ExcelStore(self.tmp, backend="sqlite")
        store.close()
        cfg = self.data / WORKSPACE_CONFIG_FILE
        data = json.loads(cfg.read_text("utf-8"))
        data.pop("storage_backend")
        cfg.write_text(json.dumps(data, ensure_ascii=False), "utf-8")
        store = ExcelStore(self.tmp)
        try:
            self.assertEqual(store.storage_backend_name, "sqlite")
            self.assertEqual(store.config.get("storage_backend"), "sqlite")
            self.assertEqual(json.loads(cfg.read_text("utf-8")).get("storage_backend"), "sqlite")
            self.assertEqual(store.config.get("data_schema_version"), "1.2.0")
        finally:
            store.close()

    def test_read_only_open_of_sqlite_workspace_has_zero_side_effects(self):
        store = ExcelStore(self.tmp, backend="sqlite")
        v = store.create_specimen()
        store.close()
        before = sorted((p.name, p.stat().st_size) for p in self.data.iterdir() if p.is_file() and not p.name.endswith("-wal") and not p.name.endswith("-shm"))
        ro = ExcelStore(self.tmp, read_only=True)
        try:
            self.assertEqual(ro.storage_backend_name, "sqlite")
            self.assertEqual(ro.get_specimen(v)["入库编号*"], v)
        finally:
            ro.close()
        after = sorted((p.name, p.stat().st_size) for p in self.data.iterdir() if p.is_file() and not p.name.endswith("-wal") and not p.name.endswith("-shm"))
        self.assertEqual(before, after)

    def test_sqlite_mode_still_writes_unmanaged_action_log_fallback_as_xlsx(self):
        # 操作记录.xlsx 不受管：sqlite 模式下仍是 xlsx 文件（Tier B 兜底路径）
        store = ExcelStore(self.tmp, backend="sqlite")
        try:
            self.assertTrue((self.data / "操作记录.xlsx").exists())
        finally:
            store.close()

    def test_schema_guard_blocks_older_software_on_sqlite_workspace(self):
        from unittest.mock import patch
        from specimen_app import excel_store as es
        from specimen_app.models import ImportConflictError

        store = ExcelStore(self.tmp, backend="sqlite")
        store.close()
        with patch.object(es, "CURRENT_DATA_SCHEMA_VERSION", "1.1.3"):
            with self.assertRaises(ImportConflictError):
                ExcelStore(self.tmp)

    def test_undo_redo_work_in_sqlite_mode(self):
        store = ExcelStore(self.tmp, backend="sqlite")
        try:
            v = store.create_specimen()
            store.set_fields("specimen", v, {"备注": "一"})
            store.set_fields("specimen", v, {"备注": "二"})
            store.undo_last()
            self.assertEqual(store.get_specimen(v)["备注"], "一")
            store.redo_last()
            self.assertEqual(store.get_specimen(v)["备注"], "二")
        finally:
            store.close()


class ConvertTests(unittest.TestCase):
    """Task 6：旧工作区（只有 xlsx）自动转换为 SQLite。"""

    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())
        self.data = self.tmp / "数据"
        from PIL import Image

        store = ExcelStore(self.tmp, backend="xlsx")
        self.v1 = store.create_specimen()
        store.set_fields("specimen", self.v1, {"管内编号*": "QD-CK-SC008-260827", "备注": "转换前"})
        store.set_fields("classification", self.v1, {"种名": "Perinereis aibuhitensis"})
        img = self.tmp / "外部" / "QD-CK-SC008-1.jpg"
        img.parent.mkdir()
        Image.new("RGB", (8, 8), "red").save(img, "JPEG")
        store.add_photo(self.v1, img, allow_outside=True)
        self.v2 = store.create_specimen()
        store.batch_reserve_vouchers(2)
        store.close()
        from specimen_app.table_backend import XlsxBackend

        self.xlsx = XlsxBackend(self.data, fit_row=lambda r, h: {k: str(r.get(k, "") or "") for k in h}, to_string=lambda v: "" if v is None else str(v))

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def _xlsx_fingerprint(self):
        # 数据版本记录.xlsx 除外：转换前的「自动快照」本身就会往里追加一行（和今天任何快照一样），
        # 其余 8 本数据表必须一字不动。
        return {p.name: (p.stat().st_mtime_ns, p.stat().st_size) for p in self.data.glob("*.xlsx") if p.name != "数据版本记录.xlsx"}

    def _version_log_types(self):
        return [r.get("操作类型", "") for r in self.xlsx.read_rows("数据版本记录.xlsx")]

    def _snapshot_count(self):
        d = self.data / "数据版本"
        return len([p for p in d.iterdir() if p.is_dir()]) if d.exists() else 0

    def _open_auto(self, **kw):
        from unittest.mock import patch
        from specimen_app import excel_store as es

        with patch.object(es, "AUTO_CONVERT_XLSX_WORKSPACES", True):
            return ExcelStore(self.tmp, **kw)

    def test_convert_moves_all_tables_and_keeps_xlsx_untouched(self):
        from specimen_app.models import MANAGED_TABLE_KEYS
        from specimen_app.excel_store import COLUMN_ALIASES

        before = self._xlsx_fingerprint()
        snaps = self._snapshot_count()
        store = self._open_auto()
        try:
            self.assertEqual(store.storage_backend_name, "sqlite")
            self.assertTrue((self.data / SQLITE_DATA_FILE).exists())
            self.assertFalse((self.data / (SQLITE_DATA_FILE + ".converting")).exists())
            self.assertEqual(self._xlsx_fingerprint(), before, "源 xlsx 不得被改动")
            self.assertIn("自动转换为 SQLite", self._version_log_types())
            self.assertEqual(store.config.get("storage_backend"), "sqlite")
            self.assertEqual(store.config.get("data_schema_version"), "1.2.0")
            self.assertEqual(self._snapshot_count(), snaps + 1, "转换前先快照")
            for key in MANAGED_TABLE_KEYS:
                self.assertEqual(store._backend.read_rows(key), self.xlsx.read_rows(key), key)
                self.assertEqual(store._backend.headers(key), [COLUMN_ALIASES.get(h, h) for h in self.xlsx.headers(key)], key)
                self.assertEqual(store._backend.exported_version(key), store._backend.table_version(key), key)
            self.assertEqual(store.get_specimen(self.v1)["备注"], "转换前")
            self.assertEqual(len(store.get_photos(self.v1)), 1)
            self.assertEqual(sorted(store.list_vouchers())[:2], sorted([self.v1, self.v2]))
        finally:
            store.close(export_excel=False)

    def test_convert_is_idempotent_on_reopen(self):
        store = self._open_auto()
        store.close(export_excel=False)
        snaps = self._snapshot_count()
        store = self._open_auto()
        try:
            self.assertEqual(store.storage_backend_name, "sqlite")
            self.assertEqual(self._snapshot_count(), snaps)
        finally:
            store.close(export_excel=False)

    def test_convert_skipped_for_read_only_store(self):
        before = sorted(p.name for p in self.data.iterdir())
        store = self._open_auto(read_only=True)
        try:
            self.assertEqual(store.storage_backend_name, "xlsx")
        finally:
            store.close()
        self.assertEqual(sorted(p.name for p in self.data.iterdir()), before)

    def test_explicit_backend_xlsx_skips_conversion(self):
        store = self._open_auto(backend="xlsx")
        try:
            self.assertEqual(store.storage_backend_name, "xlsx")
            self.assertFalse((self.data / SQLITE_DATA_FILE).exists())
        finally:
            store.close()

    def test_convert_refuses_duplicate_or_blank_headers(self):
        from openpyxl import load_workbook

        path = self.data / SPECIMEN_FILE
        wb = load_workbook(path)
        wb.active.cell(row=1, column=wb.active.max_column + 1, value="备注")  # 重复列名
        wb.save(path)
        wb.close()
        before = self._xlsx_fingerprint()
        store = self._open_auto()
        try:
            self.assertEqual(store.storage_backend_name, "xlsx")
            self.assertFalse((self.data / SQLITE_DATA_FILE).exists())
            self.assertFalse((self.data / (SQLITE_DATA_FILE + ".converting")).exists())
            self.assertEqual(self._xlsx_fingerprint(), before)
            self.assertNotIn("storage_backend", store.config)
            self.assertIn("重复", store.last_conversion_report.reason)
        finally:
            store.close()

    def test_convert_failure_leaves_no_partial_files(self):
        from unittest.mock import patch
        from specimen_app.table_backend_sqlite import SqliteBackend

        before = self._xlsx_fingerprint()
        with patch.object(SqliteBackend, "replace_many", side_effect=RuntimeError("磁盘满")):
            store = self._open_auto()
        try:
            self.assertEqual(store.storage_backend_name, "xlsx")
            self.assertFalse((self.data / SQLITE_DATA_FILE).exists())
            self.assertFalse((self.data / (SQLITE_DATA_FILE + ".converting")).exists())
            self.assertEqual(self._xlsx_fingerprint(), before)
            self.assertIn("磁盘满", store.last_conversion_report.reason)
            v = store.create_specimen()  # 仍可正常工作
            self.assertEqual(store.get_specimen(v)["入库编号*"], v)
        finally:
            store.close()

    def test_undo_still_works_after_convert(self):
        store = self._open_auto()
        try:
            store.undo_last()  # 撤销「批量预留」之前的最后一步：create_specimen(v2)
            self.assertNotIn(self.v2, store.list_vouchers())
            store.redo_last()
            self.assertIn(self.v2, store.list_vouchers())
        finally:
            store.close(export_excel=False)


if __name__ == "__main__":
    unittest.main()
